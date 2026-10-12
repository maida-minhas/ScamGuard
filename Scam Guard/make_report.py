# Generates the ScamGuard project Word document in the required report format.
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\Hp\Documents\Default Project\ScamGuard - Vision and Progress Report.docx"

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(8)
style.paragraph_format.line_spacing = 1.15

for level in (1, 2):
    h = doc.styles[f"Heading {level}"]
    h.font.name = "Calibri"
    h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def para(text, italic=False, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    if align:
        p.alignment = align
    return p


def bullets(items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


# ---------------- Title ----------------
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("ScamGuard")
r.bold = True
r.font.size = Pt(36)
r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)

s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run("Project Report")
r.font.size = Pt(18)
r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

d = doc.add_paragraph()
d.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = d.add_run("October 2026")
r.font.size = Pt(11)
r.italic = True

doc.add_paragraph()

# ---------------- 1. Overview & Problem ----------------
doc.add_heading("1. Project Overview & Problem Statement", level=1)
para(
    "ScamGuard is a community-powered website where anyone can check whether a phone number, "
    "an Instagram seller, or a Facebook Marketplace listing has been reported as a scam — and "
    "share their own experience so the next person doesn't have to learn the hard way."
)
para(
    "The problem it tackles is simple but painful. Online scams are everywhere, and the "
    "warning signs are scattered: a complaint buried in a comment thread here, a warning in a "
    "group chat there. By the time you find them, the money is usually gone. People search "
    "for trustworthy information at exactly the worst moment — right before they pay a "
    "stranger — and there is no single, honest place that answers one question: \"Has anyone "
    "else been burned here?\""
)

# ---------------- 2. Proposed Solution ----------------
doc.add_heading("2. Proposed Solution", level=1)
para(
    "ScamGuard gives every phone number, social handle, and marketplace link its own page. "
    "Anyone can search an identity in seconds and see what the community has found — star "
    "reviews, scam reports, and a plain-language risk label (\"Looks OK\", \"Caution\", or "
    "\"High risk\") so you don't need to interpret anything yourself."
)
para(
    "The key idea is that evidence, not rumor, decides what gets seen. A business review "
    "stays completely private until the reviewer uploads proof of payment and verifies their "
    "email. Each person can review or report a given identity only once, which keeps fake "
    "ratings out. Meanwhile, phone numbers and handles are shown only in masked form on "
    "public pages, so nobody's personal information is exposed."
)
para(
    "On top of this, a small built-in AI reads review text to summarize what people think "
    "and suggests the right category when someone files a report — making the whole "
    "experience faster and easier to understand."
)

# ---------------- 3. Objectives & Outcomes ----------------
doc.add_heading("3. Objectives & Expected Outcomes", level=1)
para("The project was built with these goals in mind:")
bullets([
    "Let anyone check a number, handle, or listing in seconds and understand the risk without being an expert.",
    "Make reviews trustworthy by requiring proof of payment and a verified email before anything goes public.",
    "Protect people's privacy — public pages never show raw phone numbers or handles.",
    "Give the community a simple way to report scams and see the risk picture update automatically.",
    "Use AI to make reviews easier to digest (sentiment and summaries) and reporting quicker (category suggestions).",
])
para("What we expect from it:")
bullets([
    "Users pause and check before sending money, instead of trusting a profile photo.",
    "Honest sellers get a fair, evidence-backed way to be seen as trustworthy.",
    "A growing shared record of scam activity that becomes more valuable as more people contribute.",
])

# ---------------- 4. Target Users ----------------
doc.add_heading("4. Target Users", level=1)
bullets([
    "Online shoppers and marketplace buyers — the main audience: people about to pay a stranger for goods.",
    "Students and young adults — heavy users of social commerce (Instagram sellers, Facebook Marketplace) and frequent scam targets.",
    "Small online sellers — who benefit from a reputation system that rewards proof over hype.",
    "Anyone who gets a suspicious call or message — and wants a quick second opinion before replying.",
])

# ---------------- 5. Key Features ----------------
doc.add_heading("5. Key Features / Deliverables", level=1)
doc.add_heading("Core features (delivered)", level=2)
bullets([
    "Smart search — phone numbers, Instagram handles, and Facebook Marketplace URLs each resolve to a dedicated page.",
    "Risk labels — \"Looks OK\", \"Caution\", or \"High risk\", calculated from published reviews and reports.",
    "Star reviews with a proof gate — business reviews require an uploaded proof of payment plus email verification before going public.",
    "Scam reports — categorized reports (payment fraud, non-delivery, fake seller, phishing) that feed the risk label.",
    "Identity masking — numbers and handles appear only in masked form on public pages.",
    "Personal activity page — users see everything they've written and where it stands.",
])
doc.add_heading("AI features (delivered)", level=2)
bullets([
    "Sentiment analysis — every review is labelled Positive, Negative, or Neutral, with an understanding of negation (\"not good\" counts as negative).",
    "\"What people say\" insights — a one-sentence plain-English summary of all published reviews plus the most-discussed topics.",
    "Smart category suggestions — while filing a report, one click analyses your description and picks the most likely category.",
])
para(
    "The AI runs fully locally inside the app — no external services, no API keys, no cost, "
    "and it works offline. It stays an assistant: the proof and verification gates remain the "
    "final authority over what gets published."
)
doc.add_heading("Quality deliverables", level=2)
bullets([
    "176 automated tests, all passing — covering trust gates, privacy rules, security edge cases, AI behaviour, and risk labels.",
    "Seeded demo data and a README, so anyone can run the project locally in under a minute.",
])

# ---------------- 6. Plan & Timeline ----------------
doc.add_heading("6. Implementation Plan & Timeline", level=1)
para(
    "The project was delivered in small, reviewed steps — each feature specified first, "
    "built with test-driven development, and checked before moving on."
)
bullets([
    "Phase 1 — Foundation: database schema, user accounts, and the demo email-verification flow. (Complete)",
    "Phase 2 — Search & entity pages: normalization, masking, risk labels. (Complete)",
    "Phase 3 — Trust system: star reviews, proof uploads, scam reports, activity page. (Complete)",
    "Phase 4 — AI insights: sentiment, review summaries, category suggestions. (Complete)",
    "Phase 5 — Real-world deployment: hosting, a real email service, and community growth. (Planned)",
    "Phase 6 — Reach: mobile app with call warnings, browser extension with site notifications. (Planned)",
    "Phase 7 — Escalation: option to forward confirmed scam reports to the relevant authorities. (Planned)",
])

# ---------------- 7. Resources ----------------
doc.add_heading("7. Resources or Requirements", level=1)
para("To run the current prototype:")
bullets([
    "Python 3 with a virtual environment; dependencies are listed in requirements.txt (Flask, SQLAlchemy, Flask-Login, pytest) — all free, open-source.",
    "No API keys, no paid services — the AI is built into the code itself.",
    "Commands: install requirements, initialize the database, optionally run the seed script, then start the app and open localhost:5000.",
])
para("For the planned phases:")
bullets([
    "A hosting account for the website and API (a low-cost student tier is enough).",
    "An email service for real verification links.",
    "Apple and Google developer accounts for the mobile app, and a browser store account for the extension.",
])

# ---------------- 8. Challenges ----------------
doc.add_heading("8. Potential Challenges or Risks", level=1)
bullets([
    "Cold start — a safety tool is only as strong as its community; early growth is the biggest challenge.",
    "Fairness to sellers — false or unfair reports could harm honest businesses, so a clear dispute and moderation process is needed before public launch.",
    "Privacy — identifiers must stay masked and lookups must be handled carefully as the service grows.",
    "Data freshness — keeping scam reports updating in real time across a website, app, and extension.",
    "Platform rules — call-screening and browser-extension features must pass Apple, Google, and browser-store reviews.",
    "Scope control — as a student project, staying focused (no login sprawl, no paid infrastructure) keeps it maintainable.",
])

# ---------------- 9. Additional Details ----------------
doc.add_heading("9. Additional Details", level=1)
para(
    "How to try it: run the app locally, search a seeded example such as @fashionshop.pk or a "
    "phone number, and click \"Suggest category\" on the report form to see the AI in action. "
    "Demo accounts are listed in the README; verification tokens appear on screen instead of "
    "being emailed, since this is a demo."
)
para(
    "Longer term, ScamGuard aims to become a shared neighbourhood watch for digital "
    "transactions: a website you visit today, an app that warns you when your phone rings "
    "tomorrow, and an extension that taps you on the shoulder before you pay — with a path "
    "to escalate real scams to the authorities when they are confirmed."
)

doc.save(OUT)
print(f"Saved: {OUT}")
